"""수동 상품 큐: data/상품목록.xlsx 에 플랫폼별 시트(네이버/쿠팡/기타)로 상품을 붙여넣으면 main.py --manual 이 한 줄씩 영상으로 만든다.

    python manual_queue.py --init        # 양식(엑셀)과 미디어 폴더를 만든다 (이미 있으면 건드리지 않음)
    python main.py --manual              # 모든 시트의 대기 중인 첫 줄
    python main.py --manual naver        # 네이버 시트만 (coupang, etc 도 가능)

시트 열: 상태 | 상품명 | 특징 | 제휴 링크 | 상품 페이지 | 이미지 폴더 | 결과 URL | 메모
- 상태가 비어 있거나 '대기'이고 상품명과 제휴 링크가 있는 첫 줄을 처리한다. 처리 후 '완료'(업로드) 또는 '렌더링됨'으로 바뀐다.
- 소재 폴더의 선택 파일: _대본.txt(8줄, 직접 쓴 대본으로 통째 교체) _제목.txt(상단 제목) _강조.txt(노란 강조 단어, 줄마다 하나)
  _사용.txt(사용할 이미지와 순서) _면책문구.txt(영상/설명/댓글에 들어가는 면책 문구. 네이버는 없으면 자동 생성) _추가안내.txt(설명/댓글에만 들어가는 안내, 예: 혜택 기간)
- 이미지/영상은 data/수동상품_미디어/<시트 이름>/<줄번호>/ 에 넣는다 (jpg/png/webp 와 mp4). '이미지 폴더'에 다른 경로를 적어도 된다.
"""
import argparse
from pathlib import Path

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "data" / "상품목록.xlsx"
MEDIA_ROOT = ROOT / "data" / "수동상품_미디어"
HEADERS = ["상태", "상품명", "특징", "제휴 링크", "상품 페이지", "이미지 폴더", "결과 URL", "메모"]
SHEETS = ["네이버", "쿠팡", "기타"]
ALIASES = {"naver": "네이버", "coupang": "쿠팡", "etc": "기타", "네이버": "네이버", "쿠팡": "쿠팡", "기타": "기타"}
STATES = ["대기", "완료", "렌더링됨", "건너뜀"]
IMG_EXT = {".jpg", ".jpeg", ".png", ".webp"}

# 플랫폼별 대가성 문구(영상 안에 작게 + 설명란/댓글)와 해시태그/태그
DISCLOSURES = {
    "쿠팡": ("쿠팡 파트너스 활동의 일환으로 일정액의 수수료를 제공받습니다",
             "이 포스팅은 쿠팡 파트너스 활동의 일환으로, 이에 따른 일정액의 수수료를 제공받습니다."),
    "네이버": ("네이버 쇼핑 커넥트 활동의 일환으로 판매 발생 시 수수료를 제공받습니다",
              "이 포스팅은 네이버 쇼핑 커넥트 활동의 일환으로, 판매 발생 시 수수료를 제공받습니다."),
    "알리": ("제휴 마케팅 활동의 일환으로 수수료를 받을 수 있습니다",
            "이 게시물은 제휴 마케팅 활동의 일환으로, 구매 시 일정 수수료를 지급받을 수 있습니다."),
}
DEFAULT_DISCLOSURE = DISCLOSURES["알리"]
PLATFORM_TAGS = {
    "쿠팡": (["쿠팡", "쿠팡꿀템", "꿀템추천"], ["#쿠팡", "#쿠팡파트너스"]),
    "네이버": (["네이버쇼핑", "쇼핑커넥트", "꿀템추천"], ["#네이버쇼핑", "#쇼핑커넥트"]),
    "알리": (["알리익스프레스", "알리꿀템", "꿀템추천"], ["#알리익스프레스"]),
}


def guess_brand(name: str) -> str:
    """상품명에서 브랜드를 추정: [브랜드] / 【브랜드】 가 있으면 그것, 아니면 첫 단어."""
    import re
    m = re.match(r"\s*[\[【(]([^\]】)]{1,20})[\]】)]", name)
    if m:
        return m.group(1).strip()
    first = name.strip().split()[0] if name.strip() else "해당"
    return re.sub(r"[^0-9A-Za-z가-힣&]", "", first) or "해당"


def default_brand_notice(name: str) -> str:
    """네이버 쇼핑 커넥트 영상에 항상 붙이는 문구 (브랜드 공식 입장과 다른 개인 콘텐츠/개인 의견임을 명시)."""
    return f"본 영상은 {guess_brand(name)} 공식 브랜드와 무관한 개인 콘텐츠이며 내용은 개인 의견입니다"


def disclosure_for(platform: str) -> tuple[str, str]:
    """(영상 안 짧은 문구, 설명란/댓글 문구)"""
    return DISCLOSURES.get(platform, DEFAULT_DISCLOSURE)


def init_template() -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    XLSX.parent.mkdir(exist_ok=True)
    MEDIA_ROOT.mkdir(exist_ok=True)
    if XLSX.exists():
        return XLSX
    wb = Workbook()
    wb.remove(wb.active)
    examples = {
        "네이버": ("(예시) 접이식 노트북 거치대 알루미늄", "https://naver.me/...", "https://smartstore.naver.com/..."),
        "쿠팡": ("(예시) 접이식 노트북 거치대 알루미늄", "https://link.coupang.com/a/...", "https://www.coupang.com/vp/products/..."),
        "기타": ("(예시) 접이식 노트북 거치대 알루미늄", "https://...", "https://..."),
    }
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        n, link, page = examples[name]
        ws.append(["예시", n, "알루미늄, 6단 높이조절, 접어서 휴대 (선택)", link, page, "", "", "이 줄은 '예시'라 처리되지 않습니다"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="3B2F8F")
            c.alignment = Alignment(horizontal="center")
        for col, w in zip("ABCDEFGH", (10, 38, 40, 34, 34, 28, 34, 30)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        dv = DataValidation(type="list", formula1='"' + ",".join(["예시"] + STATES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("A2:A500")
    wb.save(XLSX)
    return XLSX


def _find_media_dir(sheet: str, row: int, custom: str, name: str) -> Path:
    """소재 폴더 찾기: ① '이미지 폴더' 열 ② <시트>/<줄번호> ③ <시트>/<상품명과 같거나 앞부분이 같은 폴더>"""
    if custom:
        p = Path(custom)
        for cand in (p, MEDIA_ROOT / sheet / custom, MEDIA_ROOT / custom):     # 폴더 이름만 적어도, 절대 경로를 적어도 된다
            if cand.exists():
                return cand
        return p
    base = MEDIA_ROOT / sheet
    by_row = base / str(row)
    if by_row.exists() and any(by_row.iterdir()):
        return by_row
    if base.exists():
        key = name.replace(" ", "")[:20]
        for d in sorted(p for p in base.iterdir() if p.is_dir()):
            dn = d.name.replace(" ", "")
            if dn == name.replace(" ", "") or (key and dn.startswith(key)):
                return d
    return by_row


def next_pending(platform: str | None = None, row: int | None = None) -> dict | None:
    """상태가 비어 있거나 '대기'이고 상품명/제휴 링크가 있는 첫 줄 (platform 이 없으면 시트 순서대로 전체)."""
    from openpyxl import load_workbook
    if not XLSX.exists():
        init_template()
    wb = load_workbook(XLSX)
    want = ALIASES.get(platform or "", None) if platform and platform != "*" else None
    if platform and platform != "*" and not want:
        raise SystemExit(f"[오류] 알 수 없는 플랫폼: {platform} (naver / coupang / etc)")
    for ws in wb:
        if want and ws.title != want:
            continue
        for i, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            r = dict(zip(HEADERS, [("" if v is None else str(v).strip()) for v in cells]))
            if row is not None and i != row:       # 줄 번호를 직접 지정하면 상태와 무관하게 그 줄을 쓴다
                continue
            if row is None and r["상태"] not in ("", "대기"):
                continue
            if not (r["상품명"] and r["제휴 링크"]):
                continue
            media = _find_media_dir(ws.title, i, r["이미지 폴더"], r["상품명"])
            files = sorted(media.glob("*")) if media.exists() else []
            use_file = media / "_사용.txt"            # 사용할 이미지와 순서를 직접 지정 (줄마다 파일명 또는 파일명 일부)
            if use_file.exists():
                wanted = [l.strip() for l in use_file.read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")]
                files = [next((f for f in files if w in f.name), None) for w in wanted]
                files = [f for f in files if f]
            def _lines(name: str) -> list[str]:       # 소재 폴더의 텍스트 파일(# 주석 제외)
                fp = media / name
                return [l.strip() for l in fp.read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")] if fp.exists() else []
            notice_file = media / "_면책문구.txt"       # 브랜드 허락 조건 등 추가로 명시해야 하는 문구 (첫 줄)
            notice = notice_file.read_text(encoding="utf-8").strip().splitlines()[0] if notice_file.exists() and notice_file.read_text(encoding="utf-8").strip() else ""
            return {
                "sheet": ws.title, "row": i, "platform": ws.title, "name": r["상품명"], "features": r["특징"],
                "link": r["제휴 링크"], "page": r["상품 페이지"], "media_dir": media,
                "images": [f for f in files if f.suffix.lower() in IMG_EXT], "notice": notice,
                "extra_notes": _lines("_추가안내.txt"),
                "script_lines": _lines("_대본.txt"), "title_override": (_lines("_제목.txt") or [""])[0], "keywords": _lines("_강조.txt"),
                "videos": [f for f in files if f.suffix.lower() == ".mp4"],
            }
    return None


def mark(sheet: str, row: int, state: str, url: str = "") -> bool:
    """상태/URL 기록. 엑셀이 열려 있어 저장이 막히면 경고만 하고 계속한다 (영상은 이미 만들어졌으므로)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(XLSX)
        ws = wb[sheet]
        ws.cell(row=row, column=1, value=state)
        if url:
            ws.cell(row=row, column=7, value=url)
        wb.save(XLSX)
        return True
    except PermissionError:
        print(f"  [경고] 엑셀이 열려 있어 {sheet} {row}행 상태('{state}')를 기록하지 못했습니다. 직접 입력해 주세요." + (f" URL: {url}" if url else ""))
        return False


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--init", action="store_true")
    a = ap.parse_args()
    p = init_template()
    print(f"양식: {p}\n미디어 폴더: {MEDIA_ROOT}")
